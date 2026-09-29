"""Generate the client-facing feature roadmap Word document."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def add_feature(doc, num, title, current=None, suggestion=None, bullets=None, why=None):
    doc.add_heading(f"{num}. {title}", level=2)
    if current:
        p = doc.add_paragraph()
        p.add_run("Current state: ").bold = True
        p.add_run(current)
    if suggestion:
        p = doc.add_paragraph()
        p.add_run("Suggestion: ").bold = True
        p.add_run(suggestion)
    if bullets:
        for item in bullets:
            doc.add_paragraph(item, style="List Bullet")
    if why:
        p = doc.add_paragraph()
        p.add_run("Why it matters: ").bold = True
        p.add_run(why)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, text in enumerate(headers):
        table.rows[0].cells[index].text = text
        for paragraph in table.rows[0].cells[index].paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value


def build_document() -> Document:
    doc = Document()

    title = doc.add_heading("Portfolio Tree Explorer", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph("17 Feature Suggestions for a Complete, Polished Product")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True
    subtitle.runs[0].font.size = Pt(14)

    doc.add_paragraph(
        "Portfolio Tree Explorer is a local web application for importing portfolio hierarchy files, "
        "exploring them as an interactive tree, editing and modeling rebalancing changes, running quality checks, "
        "and generating standardized output files."
    )
    doc.add_paragraph(
        "This document lists 17 suggested enhancements to make the product more complete and production-ready, "
        "organized by priority for client review."
    )

    doc.add_heading("What the App Already Does Well", level=1)
    for item in [
        "Import hierarchy files (CSV, XLSX, ODS)",
        "Interactive tree visualization with search, zoom, pan, and custom scrollbars",
        "Edit mode: rename, change currency, move nodes, drag-and-drop, delete, bulk add",
        "Undo / redo / undo all with change audit log",
        "QC report (invalid CCY, ticker length, tree depth)",
        "Cross-held analyzer (duplicate portfolios across portgroups)",
        "Generate Output download (CSV, TXT, XLSX, ODS) reflecting in-session edits",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("High Impact — Core Workflow (Steps 1–5)", level=1)

    add_feature(
        doc,
        1,
        "Save / Round-Trip Export to Input Format",
        current=(
            "Generate Output downloads the downstream output schema (parent, portfolio_name, "
            "portfolio_type, etc.), not the original import format (Level 1, Level 2, CCY, …)."
        ),
        suggestion=(
            "Add Export hierarchy or Save as input file to export the edited tree back to the original "
            "Level 1 / Level 2 / … CSV or Excel structure so users can re-import, share with colleagues, "
            "or keep the source workbook up to date."
        ),
        why=(
            "Closes the full edit → validate → export → re-import loop for teams that work in the hierarchy template."
        ),
    )

    add_feature(
        doc,
        2,
        "Changes Only Export",
        current=(
            "The audit log and changed-node tracking exist in the browser session, but exports always "
            "include the full tree."
        ),
        bullets=[
            "Export only modified rows (input or output format)",
            "Export the audit log as CSV",
            "Show a Before vs After summary before Generate Output",
        ],
        why="Supports review and approval workflows where stakeholders need to see exactly what changed.",
    )

    add_feature(
        doc,
        3,
        "QC Gate Before Output",
        current="QC runs independently; Generate Output does not block or warn when errors exist.",
        bullets=[
            "Show a warning if QC errors are present",
            "Offer Fix first vs Export anyway",
            "Display a QC summary badge in the header (e.g. 3 errors)",
        ],
        why="Prevents accidental export of invalid portfolio data.",
    )

    add_feature(
        doc,
        4,
        "Expanded QC Rules",
        current="QC checks invalid CCY, ticker length (> 10 chars), and tree depth (> 10 levels).",
        bullets=[
            "Duplicate ticker under the same parent",
            "Empty ticker or name",
            "Invalid ticker characters",
            "Path inconsistency (Level columns vs tree path)",
            "Broken hierarchy on import",
        ],
        why="Catches more real-world data issues before export.",
    )

    add_feature(
        doc,
        5,
        "Import Validation Report",
        current="Upload succeeds or fails with a single error message.",
        bullets=[
            "Rows loaded / skipped",
            "Warnings (duplicate paths, missing CCY, unrecognized columns)",
            "Auto-run QC immediately after load",
        ],
        why="Users know upfront whether the file is clean before editing.",
    )

    doc.add_heading("Medium Impact — Polish & Usability (Steps 6–11)", level=1)

    add_feature(
        doc,
        6,
        "Tree Statistics Panel",
        bullets=["Total nodes, branches vs leaves", "Count by level and by currency", "Maximum depth"],
        why="Quick overview for large portfolios without expanding the whole tree.",
    )

    add_feature(
        doc,
        7,
        "Filter / Highlight Modes",
        bullets=[
            "Highlight by currency or level",
            "Show only branches or only leaves",
            "Show changed nodes only",
        ],
        why="Makes navigation and review faster on complex trees.",
    )

    add_feature(
        doc,
        8,
        "Keyboard Shortcuts & Navigation",
        bullets=[
            "/ — focus search",
            "Ctrl+Z / Ctrl+Y — undo / redo",
            "Arrow keys — move selection in the tree",
            "Enter — expand / collapse selected node",
        ],
        why="Power users can work much faster.",
    )

    add_feature(
        doc,
        9,
        "Session Recovery",
        bullets=[
            "Warn on page refresh if unsaved edits exist",
            "Optional localStorage backup of edited tree and audit log",
        ],
        why="Reduces lost work from accidental refresh or browser close.",
    )

    add_feature(
        doc,
        10,
        "Output Preview",
        suggestion="Before download, show a preview table with first N rows, total row count, and QC status.",
        why="Users can confirm format and content before saving the file.",
    )

    add_feature(
        doc,
        11,
        "Copy Path / Copy Row",
        bullets=[
            "Copy breadcrumb path (e.g. BCPP-ALL > BCPROP > BCPROPEQ)",
            "Copy as an input-format row for pasting into Excel",
        ],
        why="Speeds up manual updates and communication with other teams.",
    )

    doc.add_heading("Nice to Have — Completeness (Steps 12–17)", level=1)

    add_feature(
        doc,
        12,
        "Compare Two Datasets",
        suggestion=(
            "Load a baseline file vs current tree and highlight added, moved, deleted, or renamed nodes."
        ),
        why="Useful for regression checks after rebalancing or version comparisons.",
    )

    add_feature(
        doc,
        13,
        "Named Snapshots (What-If Scenarios)",
        suggestion=(
            "Save tree state as named scenarios (e.g. Scenario A, Scenario B), switch between them, and compare."
        ),
        why="Supports rebalancing exploration without losing the original structure.",
    )

    add_feature(
        doc,
        14,
        "Export Tree as Image",
        suggestion="Export the visible tree as PNG or SVG for reports and presentations.",
        why="Non-technical stakeholders often need visuals, not spreadsheets.",
    )

    add_feature(
        doc,
        15,
        "Bulk Operations",
        bullets=[
            "Bulk CCY update under a selected branch",
            "Find & replace ticker or name",
            "Bulk delete under a node (with confirmation)",
        ],
        why="Large restructuring tasks become practical.",
    )

    add_feature(
        doc,
        16,
        "README & Onboarding",
        current="README describes an earlier version of the app (basic load and drag/drop only).",
        bullets=[
            "Update README with all current features",
            "Add a short in-app tour or help panel for first-time users",
        ],
        why="Reduces learning curve and support questions.",
    )

    add_feature(
        doc,
        17,
        "Tests & Offline Reliability",
        bullets=[
            "Unit tests for portfolio_loader.py, portfolio_exporter.py, and QC rules",
            "Vendor D3.js locally (app currently loads D3 from a CDN)",
        ],
        why="Safer refactors and reliable use without internet on first load.",
    )

    doc.add_heading("Recommended Priority Order", level=1)
    add_table(
        doc,
        ["Priority", "Step", "Feature", "Main benefit"],
        [
            ("1", "1", "Round-trip input export", "Complete edit workflow"),
            ("2", "3", "QC gate before output", "Prevent bad exports"),
            ("3", "2", "Changes-only export + audit export", "Review & approval"),
            ("4", "4", "Expanded QC rules", "Catch more data issues"),
            ("5", "5", "Import validation report", "Trust at upload time"),
            ("6", "6", "Tree statistics", "Large-tree overview"),
            ("7", "7", "Filters & highlights", "Faster navigation"),
            ("8", "10", "Output preview", "Fewer download mistakes"),
            ("9", "8", "Keyboard shortcuts", "Power-user efficiency"),
            ("10", "16", "README & onboarding", "Easier adoption"),
            ("11", "9", "Session recovery", "Less lost work"),
            ("12", "11", "Copy path / row", "Manual workflow support"),
            ("13", "12", "Compare datasets", "Version control"),
            ("14", "15", "Bulk operations", "Large restructures"),
            ("15", "13", "Named snapshots", "What-if analysis"),
            ("16", "14", "Export as image", "Reporting"),
            ("17", "17", "Tests & offline D3", "Stability & offline use"),
        ],
    )

    doc.add_heading("Clarification: Generate Output vs Round-Trip Export", level=1)
    add_table(
        doc,
        ["Question", "Answer"],
        [
            ("Can users download after editing?", "Yes — via Generate Output"),
            ("Can users download as Excel?", "Yes — XLSX and ODS"),
            ("Do edits appear in the download?", "Yes — built from current in-memory tree"),
            ("Is it the same format as the uploaded file?", "No — output uses a different downstream schema"),
            (
                "Can the download be re-imported as hierarchy input?",
                "Not reliably — different column structure",
            ),
        ],
    )

    doc.add_paragraph(
        "If the workflow ends with the output file for another system, Generate Output may be sufficient. "
        "If teams also need the original hierarchy workbook updated, Step 1 remains the highest-value addition."
    )

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(
        "Portfolio Tree Explorer already delivers a strong import → explore → edit → analyze → export flow. "
        "The largest gap for a complete product is exporting back to the original hierarchy format, followed by "
        "stronger validation, change tracking in exports, and usability polish for large trees."
    )
    doc.add_paragraph(
        "Implementing these 17 steps in priority order would move the app from a capable internal tool to a "
        "polished, production-ready portfolio hierarchy platform."
    )

    footer = doc.add_paragraph(
        "Document prepared for client review — Portfolio Tree Explorer Feature Roadmap"
    )
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].italic = True
    footer.runs[0].font.size = Pt(10)
    footer.runs[0].font.color.rgb = RGBColor(100, 100, 100)

    return doc


def main() -> None:
    output = Path(__file__).resolve().parents[1] / "docs" / "Portfolio-Tree-Explorer-Feature-Roadmap.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    build_document().save(output)
    print(f"Created: {output}")


if __name__ == "__main__":
    main()
